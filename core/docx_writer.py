# -*- coding: utf-8 -*-
"""IR → .docx：表格结构、合并单元格、勾选框、印章图片都保留，且全部是可直接改的文字。"""
from __future__ import annotations

import io
from typing import Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from .ir import Document as IRDoc, Table, Cell, Para, ImageBlock

# 法院文书习惯：正文仿宋、标题黑体；这里按原 PDF 字体名映射到系统常见字体
FONT_MAP = {
    "fangsong": "仿宋", "simsun": "宋体", "simsun regular": "宋体",
    "songti": "宋体", "heiti": "黑体", "simhei": "黑体", "kaiti": "楷体",
    "microsoft yahei": "微软雅黑", "timesnewromanpsmt": "Times New Roman",
    "times new roman": "Times New Roman", "arial": "Arial",
    "fzxbsjw--gb1-0": "方正小标宋简体",
}
DEFAULT_BODY_FONT = "仿宋"
DEFAULT_TITLE_FONT = "黑体"


def map_font(name: str, fallback: str = DEFAULT_BODY_FONT) -> str:
    if not name:
        return fallback
    key = name.lower().replace(" ", "")
    for k, v in FONT_MAP.items():
        if k.replace(" ", "") == key or k.replace(" ", "") in key:
            return v
    # 中文字体名直接可用
    if any("\u4e00" <= c <= "\u9fff" for c in name):
        return name
    return fallback


def _set_run_font(run, font_name: str, size_pt: float, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)


def _set_cell_margins(cell, left=40, right=40, top=20, bottom=20) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = tc_pr.makeelement(qn("w:tcMar"), {})
        tc_pr.append(mar)
    for tag, val in (("w:top", top), ("w:left", left), ("w:bottom", bottom), ("w:right", right)):
        el = mar.find(qn(tag))
        if el is None:
            el = mar.makeelement(qn(tag), {})
            mar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


def _align_of(cell: Cell) -> str:
    """按原文文本相对格子的位置判断对齐（短标签居中，正文左对齐）。"""
    if not cell.lines:
        return "left"
    x0 = min(l.bbox.x0 for l in cell.lines)
    x1 = max(l.bbox.x1 for l in cell.lines)
    mid = (x0 + x1) / 2
    if abs(mid - cell.bbox.cx) <= max(4.0, cell.bbox.width * 0.08):
        return "center"
    if abs(x1 - cell.bbox.x1) <= max(3.0, cell.bbox.width * 0.04) and x0 > cell.bbox.x0 + cell.bbox.width * 0.15:
        return "right"
    return "left"


_ALIGN_MAP = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
              "right": WD_ALIGN_PARAGRAPH.RIGHT}


def _write_cell_text(cell: Cell, docx_cell) -> None:
    align = _align_of(cell)
    docx_cell.vertical_alignment = (WD_ALIGN_VERTICAL.CENTER if align == "center"
                                    else WD_ALIGN_VERTICAL.TOP)
    _set_cell_margins(docx_cell)
    first = True
    if not cell.lines:
        return
    for line in cell.lines:
        para = docx_cell.paragraphs[0] if first else docx_cell.add_paragraph()
        first = False
        pf = para.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        para.alignment = _ALIGN_MAP[align]
        runs = line.merged_runs()
        if not runs:
            continue
        for r in runs:
            run = para.add_run(r.text)
            _set_run_font(run, map_font(r.font), r.size_pt or 12.0, r.bold)
    for img in cell.images:
        para = docx_cell.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run()
        run.add_picture(io.BytesIO(img.png), width=Pt(max(8.0, img.bbox.width)))


def write_table(doc, table: Table) -> None:
    n_rows = max(1, table.n_rows)
    n_cols = max(1, table.n_cols)
    t = doc.add_table(rows=n_rows, cols=n_cols)
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    t.autofit = False

    # 先定列宽（合并前设置最稳）
    widths = table.col_widths or [table.bbox.width / n_cols] * n_cols
    for r in range(n_rows):
        for c in range(n_cols):
            cell = table.grid[r][c] if r < len(table.grid) and c < len(table.grid[r]) else None
            target = cell if cell is not None else _owner(table, r, c)
            w = sum(widths[c:c + (target.colspan if target else 1)]) if target else widths[c]
            t.cell(r, c).width = Pt(max(12.0, w))

    # 行高
    for r in range(n_rows):
        if r < len(table.row_heights) and table.row_heights[r] > 4:
            t.rows[r].height = Pt(table.row_heights[r])
            t.rows[r].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    # 合并
    for cell in table.cells():
        if cell.rowspan > 1 or cell.colspan > 1:
            a = t.cell(cell.row, cell.col)
            b = t.cell(min(n_rows - 1, cell.row + cell.rowspan - 1),
                       min(n_cols - 1, cell.col + cell.colspan - 1))
            if a is not b:
                try:
                    a.merge(b)
                except Exception:
                    pass

    # 填文字
    for cell in table.cells():
        try:
            _write_cell_text(cell, t.cell(cell.row, cell.col))
        except Exception as exc:  # 单个格子失败不能拖垮整篇
            cell.notes.append(f"写入失败：{exc}")


def _owner(table: Table, r: int, c: int) -> Optional[Cell]:
    """找到覆盖 (r,c) 的合并单元格。"""
    for cell in table.cells():
        if cell.row <= r < cell.row + cell.rowspan and cell.col <= c < cell.col + cell.colspan:
            return cell
    return None


def write_para(doc, para: Para) -> None:
    style = para.style
    for i, line in enumerate(para.lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0 if style == "title" else 2)
        pf.line_spacing = 1.0
        p.alignment = _ALIGN_MAP.get(para.align if style == "title" else "left", WD_ALIGN_PARAGRAPH.LEFT)
        runs = line.merged_runs()
        for r in runs:
            size = r.size_pt or (16.0 if style == "title" else 12.0)
            fam = DEFAULT_TITLE_FONT if style == "title" else map_font(r.font)
            bold = r.bold or style == "title"
            run = p.add_run(r.text)
            _set_run_font(run, fam, size, bold)


def write_image_para(doc, img: ImageBlock) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(io.BytesIO(img.png), width=Pt(max(8.0, img.bbox.width)))


def build_docx(ir: IRDoc, out_path: str) -> str:
    """把 IR 写成 .docx，返回路径。"""
    doc = Document()

    # 页面：按原 PDF 尺寸与内容边界设页边距
    page_w = max((p.width_pt for p in ir.pages), default=595.3)
    page_h = max((p.height_pt for p in ir.pages), default=841.9)
    sec = doc.sections[0]
    sec.page_width = Pt(page_w)
    sec.page_height = Pt(page_h)
    if page_w > page_h:
        sec.orientation = WD_ORIENT.LANDSCAPE
    all_x0 = [b.bbox.x0 for p in ir.pages for b in p.blocks if b.bbox.width > 0]
    all_x1 = [b.bbox.x1 for p in ir.pages for b in p.blocks if b.bbox.width > 0]
    all_y0 = [b.bbox.y0 for p in ir.pages for b in p.blocks if b.bbox.height > 0]
    left = min(max(20.0, min(all_x0) - 6.0), 90.0) if all_x0 else 56.7
    right = min(max(20.0, page_w - max(all_x1) - 6.0), 90.0) if all_x1 else 56.7
    top = min(max(20.0, min(all_y0) - 6.0), 90.0) if all_y0 else 56.7
    sec.left_margin, sec.right_margin, sec.top_margin = Pt(left), Pt(right), Pt(top)
    sec.bottom_margin = Pt(50.0)

    # 默认字体
    normal = doc.styles["Normal"]
    normal.font.name = DEFAULT_BODY_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_BODY_FONT)

    for pi, page in enumerate(ir.pages):
        if pi > 0:
            doc.add_page_break()
        for b in page.blocks:
            if b.kind == "table":
                write_table(doc, b)
            elif b.kind == "para":
                write_para(doc, b)
            elif b.kind == "image":
                write_image_para(doc, b)
    doc.save(out_path)
    return out_path
