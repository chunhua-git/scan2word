# -*- coding: utf-8 -*-
"""回读 .docx：确认表格/合并单元格/空填写栏/勾选框都在，且文字可编辑。"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
from docx import Document


def dump(path):
    d = Document(path)
    print(f"\n===== {os.path.basename(path)}")
    print(f"  段落 {len(d.paragraphs)} | 表格 {len(d.tables)} | 内嵌图片 {len(d.inline_shapes)}"
          f" | 节面 {d.sections[0].page_width.pt:.0f}x{d.sections[0].page_height.pt:.0f}pt"
          f" 边距 L{d.sections[0].left_margin.pt:.0f} T{d.sections[0].top_margin.pt:.0f}")
    for p in d.paragraphs:
        if p.text.strip():
            print(f"  ¶ [{p.alignment}] {p.text!r}")
    for ti, t in enumerate(d.tables):
        print(f"  [表{ti}] {len(t.rows)}行 x {len(t.columns)}列")
        for ri, row in enumerate(t.rows):
            cells = []
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen:
                    cells.append("⋯")
                    continue
                seen.add(id(c._tc))
                txt = c.text.replace("\n", "⏎")
                cells.append(txt if len(txt) <= 30 else txt[:28] + "…")
            print(f"    r{ri:<2} " + " | ".join(cells))


if __name__ == "__main__":
    for p in sys.argv[1:]:
        dump(p)
